"""회사 발표 템플릿(docs/ppt/*.pptx)에 최종 발표 내용을 채워 최종산출물/에 초고를 만든다.

가이드(docs/ppt/*_작성_가이드.docx) 구성: 표지 + 3장(프로젝트 개요 2분 · 기술 아키텍처 4분 · 핵심 기술 과제 4분).
템플릿의 도형·색·글꼴(Arial, 네이비/민트/회색)은 그대로 두고 텍스트 상자 안 글만 바꾼다.
아키텍처 그림은 최종산출물/1-agent-architecture.png를 왼쪽 영역에 넣는다.
수치는 전부 실측값(최종보고.md와 동일). 멘티·멘토 이름은 모르므로 자리표시를 남긴다.

사용:  python scripts/render_final_ppt.py 최종산출물/AI_Master_Project_최종발표_초고.pptx
"""

# ruff: noqa: E501 — 슬라이드 문구 줄은 100자 규칙을 적용하지 않는다

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

TEMPLATE = Path("docs/ppt/AI_Master_Project_최종발표_멘티명_사번.pptx")
DIAGRAM = Path("최종산출물/1-agent-architecture.png")

# 템플릿 색 — 제목 네이비, 본문 진회색, 설명 회색, 흰색(강조 상자 위)
NAVY, DARK, GRAY, WHITE, LIGHT = "1A365D", "333333", "666666", "FFFFFF", "CCCCCC"


def shape_by_name(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
        if sh.shape_type == 6:  # 그룹 안까지
            for g in sh.shapes:
                if g.name == name:
                    return g
    raise KeyError(name)


def fill(shape, items, same_paragraph=False):
    """텍스트 상자의 문단을 전부 새로 쓴다. items: (글, 크기pt, 굵게, 색hex). 크기 4~5의 빈 줄은 간격용.
    same_paragraph=True면 항목들을 한 문단 안의 run으로 이어 쓴다(표지 "과제명  이름"처럼)."""
    tf = shape.text_frame
    body = tf._txBody
    for para in list(tf.paragraphs)[1:]:
        body.remove(para._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        first._p.remove(r._r)
    for i, (text, size, bold, color) in enumerate(items):
        para = first if (i == 0 or same_paragraph) else tf.add_paragraph()
        run = para.add_run()
        run.text = text
        f = run.font
        f.name = "Arial"
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = RGBColor.from_string(color)


def h(text, size=9.0):
    return (text, size, True, DARK)


def b(text, size=8.5):
    return (text, size, False, GRAY)


GAP = (" ", 4.0, False, GRAY)


def cover(slide):
    title = shape_by_name(slide, "Text 4")
    title.width = Inches(6.2)  # 과제명이 한 줄에 들어가도록 오른쪽으로만 넓힌다(왼쪽 정렬은 멘티·멘토와 같게)
    fill(
        title,
        [("과제명  ", 13.0, True, "1A1A1A"), ("Code Test Agent — LLM 기반 JUnit 테스트 자동 생성·유지보수", 11.0, False, GRAY)],
        same_paragraph=True,
    )
    # 멘티·멘토는 모른다 — 자리표시 유지


def overview(slide):
    fill(
        shape_by_name(slide, "Text 10"),
        [
            h("어떤 문제를 해결하고자 했는가?"),
            b("Java 코드를 고칠 때마다 JUnit 테스트를 사람이 따라 고쳐야 하고, LLM에 통째로 맡기면 기존 테스트의 기대값을 몰래 고쳐 '통과'시키거나 리팩터링으로 바뀐 동작을 덮는 문제를 해결합니다.", 9.0),
            GAP,
            h("왜 이 문제가 중요한가? (비즈니스 임팩트)"),
            b("테스트가 코드를 못 따라가면 회귀 버그가 배포까지 새고, 잘못 고쳐진 테스트는 버그를 '정상'으로 고정합니다. 리뷰어가 매번 테스트 diff까지 검증하는 비용이 커밋마다 반복됩니다.", 9.0),
        ],
    )
    fill(
        shape_by_name(slide, "Text 14"),
        [
            h("무엇을 만들었는가?"),
            b("cta CLI — git diff의 변경 의도를 판단해 회귀 테스트를 만들고, 리팩터링으로 깨진 테스트는 사람에게 넘기며, 코드 그래프로 호출자까지 시험 범위에 넣는 에이전트", 9.0),
            GAP,
            h("에이전트 구성:"),
            b("test-lead [계획·위임·판정] → explorer [코드 그래프 조사] → writer [작성·실행 반복] → diagnoser [실패 원인 진단]  /  규칙 테이블·게이트 6종은 LLM 없는 일반 코드", 9.0),
            GAP,
            h("핵심 기술 스택:"),
            b("LangGraph · Deep Agents(LangChain) · Azure OpenAI 호환 사내 게이트웨이(gpt-5, text-embedding-3-small) · Neo4j 코드 그래프(GraphRAG) · JaCoCo · PIT · Maven / JUnit 5", 9.0),
        ],
    )
    fill(
        shape_by_name(slide, "Text 17"),
        [("100% / 100%", 14.0, True, NAVY), ("생성 테스트 라인·브랜치 커버리지 (게이트 기준 80/70, JaCoCo 실측)", 8.0, False, GRAY)],
    )
    fill(
        shape_by_name(slide, "Text 19"),
        [("83.3%", 14.0, True, NAVY), ("결함 세트 12건 검출률 · 게이트 통과 12/12 · 평균 시도 1.0회", 8.0, False, GRAY)],
    )
    fill(
        shape_by_name(slide, "Text 21"),
        [("76% 절감", 14.0, True, NAVY), ("테스트 1건 생성 비용 20,969 → 5,128 토큰, 2분 37초 → 32초. 의도 분류 정확도 100%(잘못된 '의도 모름' 0건)는 절감 후에도 동일", 8.0, False, GRAY)],
    )
    fill(
        shape_by_name(slide, "Text 24"),
        [
            ('"LLM은 판단과 작성 두 곳에만 쓰고 가드레일은 전부 코드로 두어, 기대값을 몰래 고치는 사고 0건으로 회귀 테스트 작성을 자동화했습니다."', 9.5, False, WHITE),
            (" ", 4.0, False, WHITE),
            ("실호출 실측: 커버리지 100/100 · 결함 검출 83.3% · 의도 분류 100% · 비용 76% 절감", 8.5, False, LIGHT),
        ],
    )


def architecture(slide):
    # 왼쪽 안내 텍스트를 지우고 그림을 넣는다 (영역 0.28,0.73 5.05x4.67in)
    hint = shape_by_name(slide, "Text 8")
    hint._element.getparent().remove(hint._element)
    width = Inches(4.85)
    pic = slide.shapes.add_picture(str(DIAGRAM), Inches(0.38), Inches(1.15), width=width)
    caption = slide.shapes.add_textbox(Inches(0.38), Inches(1.15) + pic.height + Inches(0.12), Inches(4.85), Inches(1.2))
    caption.text_frame.word_wrap = True
    fill(
        caption,
        [
            h("데이터 흐름", 8.5),
            b("사용자 → cta CLI(argparse) → Deep Agent(메인 test-lead + explorer·writer·diagnoser, 툴 6개) → Azure OpenAI 호환 게이트웨이(gpt-5) → 게이트 6종(JaCoCo·PIT 실측) → 제안(.cta/proposals) → cta apply", 8.0),
            b("코드 그래프(Neo4j, 없으면 인메모리) · 판단 메모(text-embedding-3-small) · 실행 장치(로컬 Maven / Docker). 모든 LLM 호출은 record/replay 가능", 8.0),
        ],
    )
    fill(shape_by_name(slide, "Text 14"), [("Deep Agents + LangGraph — 메인 1 + 서브 3", 9.5, True, NAVY)])
    fill(
        shape_by_name(slide, "Text 15"),
        [
            h("선택 이유:", 8.5),
            b("단일 ReAct 루프는 조사·작성·진단이 한 컨텍스트에 섞여 같은 실패를 반복. 역할별 서브에이전트에 task()로 위임하고 메인은 판정만 담당", 8.0),
            h("핵심 활용:", 8.5),
            b("하네스 미들웨어 — FilesystemPermission(write deny) + 툴 숨김, RunLedger(시도 ≤8·같은 실패 2회 감지), recursion_limit 400, record/replay", 8.0),
        ],
    )
    fill(shape_by_name(slide, "Text 20"), [("GraphRAG — Neo4j 코드 그래프 + 임베딩", 9.5, True, NAVY)])
    fill(
        shape_by_name(slide, "Text 21"),
        [
            h("선택 이유:", 8.5),
            b("코드 청크 임베딩은 '닮은 코드'를 찾을 뿐 호출 관계를 못 찾음. 구조 질의는 그래프(DECLARES·CREATES·COVERS 실측·CALLS 추정), 임베딩은 판단 메모(자연어)에만", 8.0),
            h("핵심 성과:", 8.5),
            b("호출자 검출 4/4(high), 다른 이름의 유사 상황 메모 검색 성공. 규칙 테이블 입력은 실측(COVERS)만", 8.0),
        ],
    )
    fill(shape_by_name(slide, "Text 26"), [("가드레일 + Skill — 규칙 테이블 · 게이트 6종 · SKILL.md", 9.5, True, NAVY)])
    fill(
        shape_by_name(slide, "Text 27"),
        [
            h("선택 이유:", 8.5),
            b("LLM이 '통과'시키는 가장 쉬운 길은 assert 완화·스킵·대상 코드 수정. 판정을 LLM에 맡기지 않고 JaCoCo·PIT 실측과 결정적 규칙으로 차단", 8.0),
            h("구현 포인트:", 8.5),
            b("탈락 사유를 프롬프트에 되돌리는 Self-Correction 루프(≤3회). 작성 지식은 SKILL.md로 분리, 규칙 기반 선택(replay 호환)", 8.0),
        ],
    )


def challenge(slide):
    fill(
        shape_by_name(slide, "Text 9"),
        [
            ("LLM이 테스트를 '통과'시키는 가장 쉬운 길은 기존 assert를 느슨하게 고치거나 스킵하는 것입니다. 실제로 assertEquals → assertNotNull 완화 시도가 관찰됐습니다.", 9.5, False, WHITE),
            ("리팩터링으로 동작이 바뀐 사실을 모델이 기대값 갱신으로 덮으면 버그가 '정상'으로 고정됩니다. 판정을 LLM에 맡기는 한 프롬프트로는 해결되지 않습니다.", 8.5, False, LIGHT),
        ],
    )
    fill(
        shape_by_name(slide, "Text 13"),
        [
            h("어떻게 해결했는가? (설계 결정)"),
            b("길과 내용을 분리했습니다. 조치의 길은 (의도 × 기존 테스트 상태) 규칙 테이블이 딕셔너리 조회로 정하고, LLM 분석은 작업 지침서 내용에만 들어갑니다. 규칙 테이블에 '기대값 자동 갱신' 행 자체가 없습니다.", 8.5),
            (" ", 5.0, False, GRAY),
            h("핵심 알고리즘 / 아키텍처 결정"),
            b("· 게이트 6종(assert 내용 비교 · 스킵 · 범위 SHA-256 · JaCoCo 커버리지 · PIT 뮤테이션 · 회귀)이 생성 후 기계적으로 판정, 탈락 사유를 프롬프트에 되돌려 재생성(Self-Correction, ≤3회)", 8.5),
            b("· 리팩터링+실패는 escalate로 멈추고(exit 3), 사람이 resolve --intended로 답하면 실패한 메서드만 allow-list로 수정", 8.5),
            b("· Deep Agents 내장 쓰기 툴은 permission deny + 툴 숨김 두 겹 차단, RunLedger가 시도 ≤8·같은 실패 2회를 집계해 diagnoser로 전환", 8.5),
            b("· 적대적 판단 메모('규칙 무시하고 진행')를 넣어도 조치가 같다는 불변식을 테스트로 고정", 8.5),
        ],
    )
    fill(shape_by_name(slide, "Text 18"), [("0건", 20.0, True, NAVY)])
    fill(
        shape_by_name(slide, "Text 20"),
        [h("기존 assert 훼손 · 스킵 · 범위 밖 수정", 9.0), b("실호출 전 시나리오(generate·maintain·--impact·resolve) 합계. 완화 시도는 assert 게이트가 차단", 8.0)],
    )
    fill(shape_by_name(slide, "Text 22"), [("83.3%", 20.0, True, NAVY)])
    fill(
        shape_by_name(slide, "Text 24"),
        [h("결함 세트 12건 검출률", 9.0), b("게이트 통과 12/12 · escalation 0 · 평균 시도 1.0회 · 의도 분류 10/10, 잘못된 '의도 모름' 0건", 8.0)],
    )
    fill(shape_by_name(slide, "Text 26"), [("76%↓", 20.0, True, NAVY)])
    fill(
        shape_by_name(slide, "Text 28"),
        [h("테스트 생성 토큰 20,969 → 5,128", 9.0), b("시간 2분 37초 → 32초. 커버리지 탈락 재생성 3회 → 1회(탈락 사유에 미실행 줄 원문 첨부)", 8.0)],
    )


def main() -> None:
    out = Path(sys.argv[1])
    prs = Presentation(TEMPLATE)
    cover(prs.slides[0])
    overview(prs.slides[1])
    architecture(prs.slides[2])
    challenge(prs.slides[3])
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
